"""
============================================================
app/services/export/excel_exporter.py

Excel Exporter

Responsibilities
----------------
- Export a document to XLSX
- Build workbook
- Format headers
- Export line items

No database.
No repositories.
============================================================
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from app.models.domain.document import (
    Document,
    DocumentLineItem,
)


class ExcelExporter:
    """
    Exports documents to XLSX.
    """

    HEADER_FILL = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    HEADER_FONT = Font(
        bold=True,
        color="FFFFFF",
    )

    TITLE_FONT = Font(
        bold=True,
        size=14,
    )

    # =====================================================
    # Export
    # =====================================================

    def export(
        self,
        *,
        document: Document,
        line_items: list[DocumentLineItem],
    ) -> bytes:

        workbook = Workbook()

        self._summary_sheet(
            workbook,
            document,
        )

        self._line_items_sheet(
            workbook,
            line_items,
        )

        stream = BytesIO()

        workbook.save(stream)

        stream.seek(0)

        return stream.getvalue()

    # =====================================================
    # Summary
    # =====================================================

    def _summary_sheet(
        self,
        workbook: Workbook,
        document: Document,
    ) -> None:

        sheet = workbook.active

        sheet.title = "Summary"

        sheet["A1"] = "Document Summary"
        sheet["A1"].font = self.TITLE_FONT

        rows = [
            ("Document Number", document.document_number),
            ("Vendor", document.vendor_name),
            ("Document Type", document.document_type.value),
            ("Status", document.status.value),
            ("Currency", document.currency),
            ("Subtotal", document.subtotal),
            ("Tax", document.tax_amount),
            ("Discount", document.discount_amount),
            ("Total", document.total_amount),
            ("Amount Paid", document.amount_paid),
            ("Amount Due", document.amount_due),
            ("Document Date", document.document_date),
            ("Due Date", document.due_date),
        ]

        row = 3

        for key, value in rows:

            sheet.cell(
                row=row,
                column=1,
                value=key,
            ).font = Font(
                bold=True,
            )

            sheet.cell(
                row=row,
                column=2,
                value=value,
            )

            row += 1

        sheet.column_dimensions["A"].width = 30
        sheet.column_dimensions["B"].width = 30

    # =====================================================
    # Line Items
    # =====================================================

    def _line_items_sheet(
        self,
        workbook: Workbook,
        line_items: list[DocumentLineItem],
    ) -> None:

        sheet = workbook.create_sheet(
            "Line Items",
        )

        headers = [
            "Description",
            "Quantity",
            "Unit Price",
            "Amount",
            "Tax",
            "Discount",
            "Category",
            "SKU",
        ]

        for column, header in enumerate(
            headers,
            start=1,
        ):

            cell = sheet.cell(
                row=1,
                column=column,
                value=header,
            )

            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT

        row = 2

        for item in line_items:

            sheet.cell(row=row, column=1).value = item.description
            sheet.cell(row=row, column=2).value = item.quantity
            sheet.cell(row=row, column=3).value = item.unit_price
            sheet.cell(row=row, column=4).value = item.amount
            sheet.cell(row=row, column=5).value = item.tax_amount
            sheet.cell(row=row, column=6).value = item.discount
            sheet.cell(row=row, column=7).value = item.category
            sheet.cell(row=row, column=8).value = item.sku

            row += 1

        for column in range(1, len(headers) + 1):

            sheet.column_dimensions[
                get_column_letter(column)
            ].width = 20